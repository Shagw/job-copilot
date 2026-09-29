"""Extract plain text from uploaded resumes (PDF, DOCX, TXT)."""
import io
import zipfile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class FileParseError(ValueError):
    pass


def _ext(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF"):
        raise FileParseError("File is not a valid PDF")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise FileParseError("Password-protected PDFs are not supported")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as e:
        raise FileParseError("Could not read PDF") from e


def _docx(data: bytes) -> str:
    if not data.startswith(b"PK"):
        raise FileParseError("File is not a valid DOCX")
    try:
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as e:
        raise FileParseError("Could not read DOCX") from e
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many resumes use tables for layout
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_text(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise FileParseError("Only PDF, DOCX and TXT files are supported")
    text = {".pdf": _pdf, ".docx": _docx, ".txt": _txt}[ext](data)
    # Normalize whitespace but keep line structure (bullets matter for chunking).
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    text = "\n".join(lines).strip()
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    if len(text) < 50:
        raise FileParseError("Could not find enough text in this file. Is it a scanned image?")
    return text
