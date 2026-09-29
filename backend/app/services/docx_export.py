"""Plain-text resume / cover letter -> simple, ATS-friendly DOCX."""
import io
import re

from docx import Document
from docx.shared import Pt

_HEADING = re.compile(r"^[A-Z][A-Z &/,\-]{2,40}$")  # "EXPERIENCE", "SKILLS & TOOLS"


def text_to_docx(text: str) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for line in text.splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        if stripped.startswith(("- ", "• ", "* ")):
            doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        elif _HEADING.match(stripped):
            doc.add_heading(stripped, level=2)
        else:
            doc.add_paragraph(stripped)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
