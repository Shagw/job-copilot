import io

import pytest
from docx import Document
from sqlalchemy import select

from app.models import Resume
from app.rag.store import chunk_text
from app.services.file_parser import FileParseError, extract_text

RESUME = """Alice Example
Software Engineer | alice@example.com

EXPERIENCE
Acme Corp - Backend Engineer (2021-2024)
- Built REST APIs in Python with FastAPI serving 2M requests per day
- Deployed services on Docker and AWS ECS, cut deploy time by 40%
- Set up CI/CD pipelines with GitHub Actions

Globex - Data Intern (2020)
- Wrote ETL jobs in SQL and pandas for sales reporting

SKILLS
Python, FastAPI, PostgreSQL, Docker, AWS, React
"""


def make_pdf(text: str) -> bytes:
    """Minimal valid one-page PDF with the given lines of text."""
    lines = text.splitlines()
    ops = ["BT", "/F1 11 Tf", "50 780 Td", "14 TL"]
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({safe}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1", "replace")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def make_docx(text: str) -> bytes:
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Certification"
    table.rows[0].cells[1].text = "AWS Solutions Architect"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- file parsing ----------

def test_extract_txt():
    assert "FastAPI" in extract_text("cv.txt", RESUME.encode())


def test_extract_pdf():
    text = extract_text("cv.PDF", make_pdf(RESUME))
    assert "Acme Corp" in text and "GitHub Actions" in text


def test_extract_docx_includes_tables():
    text = extract_text("cv.docx", make_docx(RESUME))
    assert "Acme Corp" in text and "AWS Solutions Architect" in text


@pytest.mark.parametrize(
    "name,data",
    [
        ("cv.exe", b"MZ..."),
        ("cv.pdf", b"not really a pdf" * 10),  # wrong magic bytes
        ("cv.docx", b"not a zip" * 10),
        ("cv.docx", b"PK\x03\x04garbage" * 10),  # zip header but broken
        ("cv.txt", b"too short"),
    ],
)
def test_bad_files_rejected(name, data):
    with pytest.raises(FileParseError):
        extract_text(name, data)


# ---------- chunking ----------

def test_chunks_respect_size_and_keep_all_content():
    text = "\n\n".join(f"Section {i}\n" + "\n".join(f"- bullet {i}.{j} " + "x" * 80 for j in range(6))
                       for i in range(5))
    chunks = chunk_text(text, size=600, overlap=100)
    assert len(chunks) > 1
    assert all(len(c) <= 600 + 100 + 1 for c in chunks)
    for i in range(5):
        for j in range(6):
            assert any(f"bullet {i}.{j}" in c for c in chunks)


def test_giant_single_line_is_split():
    chunks = chunk_text("word " * 1000, size=600, overlap=100)
    assert len(chunks) > 5


# ---------- RAG index ----------

def test_search_returns_relevant_chunk(resume_index):
    resume_index.index_resume(1, 1, RESUME)
    hits = resume_index.search(1, "Docker AWS deployment")
    assert hits and any("Docker" in h.text for h in hits)


def test_search_is_isolated_per_user(resume_index):
    resume_index.index_resume(1, 1, RESUME)
    resume_index.index_resume(2, 2, "Bob Builder\n\nEXPERIENCE\n- Welding and carpentry for 10 years on bridges")
    assert all("Bob" not in h.text and "Welding" not in h.text for h in resume_index.search(1, "welding carpentry"))
    assert all("Acme" not in h.text for h in resume_index.search(2, "Python FastAPI Docker"))


def test_reindex_replaces_old_vectors(resume_index):
    resume_index.index_resume(1, 1, RESUME)
    resume_index.index_resume(1, 2, "Alice Example\n\nNow a chef\n- Cooked pasta at Luigi's restaurant for three years")
    hits = resume_index.search(1, "Python FastAPI Docker", k=10)
    assert all("Acme" not in h.text for h in hits)


def test_search_for_user_without_resume_is_empty(resume_index):
    assert resume_index.search(99, "python") == []


# ---------- endpoints ----------

def upload(client, name="cv.txt", data=RESUME.encode()):
    return client.post("/resume", files={"file": (name, data, "application/octet-stream")})


def test_upload_requires_login(client, resume_index):
    assert upload(client).status_code == 401
    assert client.get("/resume").status_code == 401


def test_upload_and_get(logged_in, resume_index):
    r = upload(logged_in, "cv.pdf", make_pdf(RESUME))
    assert r.status_code == 201
    body = r.json()
    assert body["filename"] == "cv.pdf" and body["chunks"] >= 1 and "Acme" in body["raw_text"]
    assert logged_in.get("/resume").json()["id"] == body["id"]


def test_get_without_upload_is_404(logged_in, resume_index):
    assert logged_in.get("/resume").status_code == 404


def test_reupload_replaces_active_resume_but_keeps_history(logged_in, resume_index, db):
    first = upload(logged_in).json()
    second = upload(logged_in, data=(RESUME + "\nNEW: Kubernetes certified").encode()).json()
    assert logged_in.get("/resume").json()["id"] == second["id"] != first["id"]
    assert len(db.scalars(select(Resume)).all()) == 2  # nothing deleted


def test_upload_too_large(logged_in, resume_index):
    r = upload(logged_in, data=b"a" * (5 * 1024 * 1024 + 1))
    assert r.status_code == 413


def test_upload_bad_type(logged_in, resume_index):
    r = upload(logged_in, "cv.exe", b"MZ" * 100)
    assert r.status_code == 400 and "PDF, DOCX and TXT" in r.json()["detail"]


def test_filename_path_is_stripped(logged_in, resume_index):
    r = upload(logged_in, "../../etc/cv.txt")
    assert r.json()["filename"] == "cv.txt"


def test_users_cannot_see_each_others_resume(client, outbox, resume_index):
    from tests.conftest import make_user

    make_user(client, outbox, stay_logged_in=True)
    upload(client)
    client.post("/auth/logout")
    make_user(client, outbox, email="bob@example.com", stay_logged_in=True)
    assert client.get("/resume").status_code == 404
